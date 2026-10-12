"""Create a demo spreadsheet of FAKE cases with deadlines relative to today.

    python make_sample_data.py          # writes cases.xlsx
    python make_sample_data.py --csv    # writes cases.csv instead
    python make_sample_data.py --out demo/cases_demo.xlsx   # any path (used by the one-click demo)

Re-run it before a demo so every urgency level shows up.
All names are invented and all emails use the reserved example.com domain.
"""

from __future__ import annotations

import argparse
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

from tracker.data_source import COLUMNS
from tracker.settings import PROJECT_DIR

ATTORNEYS = {
    "Maria Delgado": "maria.delgado@example.com",
    "James Okafor": "james.okafor@example.com",
    "Priya Raman": "priya.raman@example.com",
}

# (client, case type, status, days from today, deadline type, attorney)
CASES = [
    ("Ana Morales", "RFE Response", "Drafting response", -6, "RFE response due", "Maria Delgado"),
    ("Wei Zhang", "I-485", "Awaiting client documents", -2, "Filing deadline", "James Okafor"),
    ("Olusegun Adeyemi", "Asylum", "Preparing evidence packet", -1, "Filing deadline", "Priya Raman"),
    ("Sofia Petrova", "I-130", "Ready for attorney review", 0, "Filing deadline", "Maria Delgado"),
    ("Rahul Mehta", "N-400", "Biometrics scheduled", 1, "Biometrics appointment", "James Okafor"),
    ("Fatima Haddad", "Asylum", "Hearing prep", 2, "Court hearing", "Priya Raman"),
    ("Carlos Ramirez", "I-485", "Interview scheduled", 4, "Biometrics appointment", "Maria Delgado"),
    ("Linh Nguyen", "RFE Response", "Gathering evidence", 6, "RFE response due", "James Okafor"),
    ("Daniel Kowalski", "I-130", "Awaiting client documents", 7, "Filing deadline", "Priya Raman"),
    ("Grace Mensah", "N-400", "Filed – awaiting interview", 12, "Biometrics appointment", "Maria Delgado"),
    ("Hiroshi Tanaka", "I-485", "Preparing forms", 21, "Filing deadline", "James Okafor"),
    ("Amara Okeke", "Asylum", "Declaration in progress", 30, "Court hearing", "Priya Raman"),
    ("Mateo Silva", "I-130", "Filed – awaiting receipt", 45, "Filing deadline", "Maria Delgado"),
    ("Yasmin Farouk", "N-400", "Interview prep", 75, "Court hearing", "James Okafor"),
    ("Elena Rossi", "I-485", "Medical exam pending", 3, "Filing deadline", "Priya Raman"),
]

# Rows with deliberate problems, to demo the "Needs attention" section.
BAD_ROWS = [
    {"Client Name": "Tomas Herrera", "Case Type": "I-130", "Current Status": "Intake",
     "Next Deadline": None, "Deadline Type": "Filing deadline",
     "Assigned Attorney": "Maria Delgado", "Attorney Email": ATTORNEYS["Maria Delgado"]},
    {"Client Name": "Nadia Rahman", "Case Type": "RFE Response", "Current Status": "Drafting response",
     "Next Deadline": "TBD", "Deadline Type": "RFE response due",
     "Assigned Attorney": "James Okafor", "Attorney Email": ATTORNEYS["James Okafor"]},
    {"Client Name": "Kofi Asante", "Case Type": "Asylum", "Current Status": "Intake",
     "Next Deadline": "+10", "Deadline Type": "Court hearing",
     "Assigned Attorney": None, "Attorney Email": None},
]


def build_rows(today: date) -> list[dict]:
    rows = []
    for n, (client, case_type, status, offset, deadline_type, attorney) in enumerate(CASES, start=1):
        rows.append({
            "Case ID": f"IMM-{today.year}-{n:03d}",
            "Client Name": client,
            "Case Type": case_type,
            "Current Status": status,
            "Next Deadline": today + timedelta(days=offset),
            "Deadline Type": deadline_type,
            "Assigned Attorney": attorney,
            "Attorney Email": ATTORNEYS[attorney],
        })
    for n, bad in enumerate(BAD_ROWS, start=len(CASES) + 1):
        row = {"Case ID": f"IMM-{today.year}-{n:03d}", **bad}
        if row["Next Deadline"] == "+10":  # A valid date, but no attorney assigned.
            row["Next Deadline"] = today + timedelta(days=10)
        rows.append(row)
    return rows


def write_xlsx(rows: list[dict], path) -> None:
    from openpyxl.styles import Font, PatternFill

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=list(COLUMNS)).to_excel(writer, index=False, sheet_name="Cases")
        ws = writer.sheets["Cases"]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
        for col, width in zip("ABCDEFGH", (15, 20, 14, 28, 14, 24, 18, 28)):
            ws.column_dimensions[col].width = width
        for cell in ws["E"][1:]:
            cell.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--csv", action="store_true", help="Write cases.csv instead of cases.xlsx")
    parser.add_argument("--out", help="Write to this .xlsx/.csv path instead (keeps real case files untouched)")
    args = parser.parse_args()

    rows = build_rows(date.today())
    path = Path(args.out).expanduser() if args.out else PROJECT_DIR / ("cases.csv" if args.csv else "cases.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        df = pd.DataFrame(rows, columns=list(COLUMNS))
        df["Next Deadline"] = df["Next Deadline"].map(lambda d: d.isoformat() if isinstance(d, date) else (d or ""))
        df.to_csv(path, index=False)
    else:
        write_xlsx(rows, path)
    print(f"Wrote {len(rows)} fake cases to {path}")


if __name__ == "__main__":
    main()
